# -*- coding: utf-8 -*-
"""详细设计说明书：改标题、填成员、插入两张状态机图。
编码规范自查表：清空问题说明列与整改清单。"""
from docx import Document
from docx.shared import Inches

DETAIL = r"D:\桌面\code\javafx\bilitro\XX项目_详细设计说明书_v3.0_已填写.docx"
CHECK = r"D:\桌面\code\javafx\bilitro\XX项目_编码规范自查表_v3.0_已填写.docx"

# ===== 详细设计说明书 =====
d = Document(DETAIL)
for p in d.paragraphs:
    t = p.text.strip()
    if t == "需求规格说明书":
        for r in p.runs:
            r.text = ""
        p.runs[0].text = "详细设计说明书"
    elif t.startswith("小组成员"):
        for r in p.runs:
            r.text = ""
        p.runs[0].text = "小组成员：苏鹏隐 陈科宇 方腾 黄思贤 张荣泽"
    elif t == "在此插入状态机图":
        for r in p.runs:
            r.text = ""
        p.runs[0].add_picture(r"D:\桌面\code\javafx\bilitro\state_main.png", width=Inches(6.2))
    elif t == "在此插入子状态机图":
        for r in p.runs:
            r.text = ""
        p.runs[0].add_picture(r"D:\桌面\code\javafx\bilitro\state_sub.png", width=Inches(6.2))
d.save(DETAIL)
print("detail saved")

# ===== 编码规范自查表：清空问题说明列与整改清单 =====
c = Document(CHECK)
for ti in (0, 2, 4, 6, 8, 10):  # 六个维度检查表
    t = c.tables[ti]
    for row in t.rows[1:]:
        row.cells[3].text = ""
for ti in (1, 3, 5, 7, 9, 11):  # 整改清单
    t = c.tables[ti]
    for row in t.rows[1:]:
        for cell in row.cells:
            cell.text = ""
c.save(CHECK)
print("check saved")
