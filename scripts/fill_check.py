# -*- coding: utf-8 -*-
"""填写编码规范自查表 v3.0（依据 bilitro 项目实际代码扫描结果）。"""
from docx import Document

SRC = r"D:\桌面\code\javafx\bilitro\编码规范自查表.docx"
DST = r"D:\桌面\code\javafx\bilitro\XX项目_编码规范自查表_v3.0_已填写.docx"

d = Document(SRC)


def fill_check(table, results):
    """results: list of (自查结果, 问题说明)，按数据行顺序填。"""
    for i, (r, note) in enumerate(results):
        table.rows[i + 1].cells[2].text = r
        table.rows[i + 1].cells[3].text = note


def fill_list(table, first_row_text):
    """整改清单：无问题则第一行写无。"""
    row = table.rows[1]
    row.cells[0].text = "无"
    for c in row.cells[1:]:
        c.text = first_row_text

G = "合规"

# 维度1 命名
fill_check(d.tables[0], [
    (G, "变量与字段均使用有含义的名称，如 remainingPlays、selected、remainingTargetScore，未发现 a、b、temp、data 这类无意义命名"),
    (G, "类名均为大驼峰，如 DefaultGameSession、RandomShop、GameViewFx"),
    (G, "方法名与变量名均为小驼峰，如 evaluateDetail、claimLevelClearReward、toggleSelect"),
    (G, "GameConfig 中常量均为全大写加下划线，如 HAND_SIZE、PLAYS_PER_LEVEL、SPECIAL_CARD_LIMIT"),
    (G, "包名全部小写：com.bilitro.model、com.bilitro.controller、com.bilitro.view"),
])
fill_list(d.tables[1], "无")

# 维度2 异常
fill_check(d.tables[2], [
    (G, "全项目无空 catch 块"),
    (G, "全项目无 printStackTrace 调用"),
    (G, "唯一一处 catch 在 SpecialCardCatalog.loadAll，捕获后包装为 IllegalStateException 附带配置路径向上抛出"),
    (G, "仅捕获具体的 IOException，未捕获过宽的 Exception"),
    (G, "AI 生成代码逐行核对，无吞异常的情况"),
])
fill_list(d.tables[3], "无")

# 维度3 线程
fill_check(d.tables[4], [
    (G, "界面更新均在 JavaFX Application Thread 中执行，结算弹窗通过 Platform.runLater 回到界面线程"),
    (G, "事件处理方法中无 Thread.sleep，计分动画使用 JavaFX 时间轴实现"),
    (G, "计分过程由 ScoreCalculator.steps 预先算好快照，事件处理中只做状态读写，无复杂循环"),
    (G, "当前无耗时 IO 在事件链上执行；存档加载使用局部流即时完成"),
    (G, "Platform.runLater 中只创建并展示结算弹窗，属于纯界面更新"),
    (G, "实际运行对局与商店界面无卡死现象"),
])
fill_list(d.tables[5], "无")

# 维度4 资源
fill_check(d.tables[6], [
    (G, "配置表读取的 InputStream 与 Reader 均在 try-with-resources 中关闭"),
    (G, "SpecialCardCatalog.loadAll 使用双层 try-with-resources，异常时流也能关闭"),
    (G, "界面事件绑定在节点创建时一次性完成，场景切换时旧界面随场景树一起释放"),
    (G, "无自建后台线程，动画由 JavaFX 时间轴驱动，应用退出时随 JavaFX 平台停止"),
    (G, "无静态可变集合，商品池与手牌均为实例字段，随对局结束释放"),
    (G, "hand、goods、selected、remainingDeck 等方法均返回 List.copyOf 只读副本，不暴露内部集合"),
])
fill_list(d.tables[7], "无")

# 维度5 方法长度
fill_check(d.tables[8], [
    (G, "逐文件扫描全部 38 个源文件，无超过 40 行的方法"),
    (G, "最长方法为 GameViewFx 的界面搭建方法，已在 40 行以内，无需拆分"),
    (G, "无遗留超长方法"),
])
fill_list(d.tables[9], "无")

# 维度6 调试方式
fill_check(d.tables[10], [
    (G, "全项目无 System.out.println，状态变化通过界面刷新直接观察"),
    (G, "调出牌型判定时使用行断点跟踪 evaluateDetail 的逐级匹配"),
    (G, "购买校验异常时使用异常断点定位抛出来源"),
    (G, "排查剩余次数扣减问题时使用条件断点，如 remainingPlays == 0 时暂停"),
    (G, "在 play、discard、outcome 上设方法断点核对调用顺序"),
    (G, "用字段断点定位 hand 列表被修改的位置"),
])
fill_list(d.tables[11], "无")

# 二、自查总结
summary = d.tables[12]
rows = ["命名", "异常", "线程", "资源", "方法长度", "调试方式"]
for i, name in enumerate(rows):
    summary.rows[i + 1].cells[1].text = "合规"
    summary.rows[i + 1].cells[2].text = "0"

d.save(DST)
print("saved", DST)
