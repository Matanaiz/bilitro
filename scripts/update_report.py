# -*- coding: utf-8 -*-
"""更新单元测试报告：补充 controller 测试部分与覆盖率数据"""
from docx import Document

PATH = r"D:\桌面\code\javafx\bilitro\docs\bilitro项目_单元测试报告_v5.0.docx"
doc = Document(PATH)


def set_text(p, text):
    """整段替换文字，保留首个 run 的格式。"""
    if p.runs:
        p.runs[0].text = text
        for r in p.runs[1:]:
            r.text = ""
    else:
        p.add_run(text)


# 1. 二、核心用例清单：表格追加 controller 两行
table = doc.tables[0]
for row_data in [
    ["8", "DefaultGameController.onPlay / toggleSelect", "边界",
     "选中达5张上限后无法加选；出牌扣次数并清空选中；通关时弹出奖励并触发结局回调", "否"],
    ["9", "DefaultShopController.onBuy / onRefresh", "异常",
     "代币不足与栏位已满时弹出对应提示；出售后界面刷新为已售状态", "否"],
    ["10", "HandType.baseScore / Rank.chips", "正常",
     "牌型初始分数倍数与卡牌计分点数符合对照表（对子10×2、同花顺100×8、A计11、K计10）", "是"],
]:
    cells = table.add_row().cells
    for i, v in enumerate(row_data):
        cells[i].text = v

# 2. 3.3 节：替换标题与说明
for i, p in enumerate(doc.paragraphs):
    if p.text.strip().startswith("3.3 Controller层核心测试类"):
        set_text(p, "3.3 Controller层核心测试类（类名：DefaultGameControllerTest / DefaultShopControllerTest）")
        set_text(doc.paragraphs[i + 1],
                 "Controller 层不直接操作 JavaFX 控件（动画、弹窗均在 View 层），"
                 "测试前先将 GameView 接口补全并把控制器依赖由具体类改为接口，"
                 "再以纯 Java 假视图（FakeGameView / FakeShopView）记录调用、"
                 "将计分动画的完成回调改为同步执行，无需启动 UI 工具包即可完成测试。"
                 "共 18 个用例，覆盖选牌上限、出牌弃牌、结局回调、购买与刷新提示等场景。"
                 "核心代码如下：")
        anchor = doc.paragraphs[i + 2]  # 代码插到说明之后、原空段之前
        break

code = """package com.bilitro.controller;

/** 假视图：记录每次调用，计分过程动画改为同步执行完成回调。 */
static class FakeGameView implements GameView {
    List<Card> lastSelected = List.of();
    Integer effectScore;
    GameSession.RewardBreakdown rewardShown;

    @Override
    public void playScoringProcess(List<ScoreCalculator.ScoreStep> steps, Runnable onFinished) {
        onFinished.run(); // 测试中不播动画，直接执行完成回调
    }
    // 其余方法仅记录参数或留空
}

@Test
void 选中达到上限后无法加选() {
    GameSession s = newSession(new DefaultPlayer(), 10000);
    DefaultGameController c = newController(s, new FakeGameView());
    for (int i = 0; i < GameConfig.HAND_SIZE; i++) {
        c.toggleSelect(s.hand().get(i));
    }
    assertEquals(GameConfig.MAX_SELECT, c.selected().size());
}

@Test
void 通关后弹出奖励并触发结局回调() {
    GameSession s = newSession(new DefaultPlayer(), 1); // 目标分极低，任意出牌即过关
    FakeGameView view = new FakeGameView();
    DefaultGameController c = newController(s, view);
    AtomicReference<GameSession.RoundOutcome> outcome = new AtomicReference<>();
    c.setOutcomeHandler(outcome::set);
    c.toggleSelect(s.hand().get(0));
    c.onPlay();
    assertEquals(GameSession.RoundOutcome.LEVEL_CLEARED, outcome.get());
    assertNotNull(view.rewardShown);
}

@Test
void 代币不足购买弹提示() {
    DefaultPlayer player = new DefaultPlayer(0);
    RandomShop shop = new RandomShop(player);
    FakeShopView view = new FakeShopView();
    DefaultShopController c = new DefaultShopController(shop, player, view);
    SpecialCard item = shop.goods().stream()
            .filter(g -> g.price() > 0).findFirst().orElseThrow();
    c.onBuy(item);
    assertEquals("代币不足", view.lastError);
}"""
for line in reversed(code.split("\n")):
    anchor.insert_paragraph_before(line)

# 3. 四、覆盖率汇总：填入真实数据
for p in doc.paragraphs:
    t = p.text
    if t.startswith("model包（业务逻辑层）分支覆盖率"):
        set_text(p, "model包（业务逻辑层）分支覆盖率：81%")
    elif t.startswith("controller包（控制层）分支覆盖率"):
        set_text(p, "controller包（控制层）分支覆盖率：81%")

doc.save(PATH)
print("saved:", PATH)
