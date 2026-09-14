# -*- coding: utf-8 -*-
"""修复 3.3 节代码段顺序：删除插反的段落，按正确顺序重新插入"""
from docx import Document

PATH = r"D:\桌面\code\javafx\bilitro\docs\bilitro项目_单元测试报告_v5.0.docx"
doc = Document(PATH)

# 定位说明段（3.3 标题的下一段）
exp_idx = None
for i, p in enumerate(doc.paragraphs):
    if p.text.startswith("Controller 层不直接操作 JavaFX 控件"):
        exp_idx = i
        break
assert exp_idx is not None

# 说明段之后、下一个 Heading 之前的段落全部视为插反的代码块，删除
head_idx = None
for j in range(exp_idx + 1, len(doc.paragraphs)):
    if doc.paragraphs[j].style.name.startswith("Heading"):
        head_idx = j
        break
assert head_idx is not None
for p in doc.paragraphs[exp_idx + 1:head_idx]:
    p._element.getparent().remove(p._element)

anchor = doc.paragraphs[exp_idx + 1]  # 删除后 Heading 空段紧跟说明段

code = """package com.bilitro.controller;

/** 假视图：记录每次调用，计分过程动画改为同步执行完成回调。 */
static class FakeGameView implements GameView {
    List<Card> lastSelected = List.of();
    Integer effectScore;
    GameSession.RewardBreakdown rewardShown;

    @Override
    public void playScoringProcess(List<ScoreCalculator.ScoreStep> steps, Runnable onFinished) {
        onFinished.run(); // 测试中不播动画，直接执行完成回调
    }
    // 其余方法仅记录参数或留空
}

@Test
void 选中达到上限后无法加选() {
    GameSession s = newSession(new DefaultPlayer(), 10000);
    DefaultGameController c = newController(s, new FakeGameView());
    for (int i = 0; i < GameConfig.HAND_SIZE; i++) {
        c.toggleSelect(s.hand().get(i));
    }
    assertEquals(GameConfig.MAX_SELECT, c.selected().size());
}

@Test
void 通关后弹出奖励并触发结局回调() {
    GameSession s = newSession(new DefaultPlayer(), 1); // 目标分极低，任意出牌即过关
    FakeGameView view = new FakeGameView();
    DefaultGameController c = newController(s, view);
    AtomicReference<GameSession.RoundOutcome> outcome = new AtomicReference<>();
    c.setOutcomeHandler(outcome::set);
    c.toggleSelect(s.hand().get(0));
    c.onPlay();
    assertEquals(GameSession.RoundOutcome.LEVEL_CLEARED, outcome.get());
    assertNotNull(view.rewardShown);
}

@Test
void 代币不足购买弹提示() {
    DefaultPlayer player = new DefaultPlayer(0);
    RandomShop shop = new RandomShop(player);
    FakeShopView view = new FakeShopView();
    DefaultShopController c = new DefaultShopController(shop, player, view);
    SpecialCard item = shop.goods().stream()
            .filter(g -> g.price() > 0).findFirst().orElseThrow();
    c.onBuy(item);
    assertEquals("代币不足", view.lastError);
}"""
for line in code.split("\n"):
    anchor.insert_paragraph_before(line)

doc.save(PATH)
print("fixed:", PATH)
