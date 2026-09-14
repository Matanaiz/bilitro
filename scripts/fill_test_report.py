# 填充单元测试报告：二、核心用例清单；三、核心测试代码
from docx import Document
from docx.shared import Pt
from docx.oxml.ns import qn

PATH = 'bilitro项目_单元测试报告_v5.0.docx'
d = Document(PATH)

# ---------- 标题信息 ----------
for p in d.paragraphs:
    if p.text.strip() == 'xx项目':
        p.text = 'bilitro项目'
    elif '小组成员' in p.text:
        p.text = '小组成员：苏鹏隐 陈科宇 方腾 黄思贤 张荣泽'
    elif 'xxxx年xx月xx日' in p.text:
        p.text = '2026年9月14日'

# ---------- 二、核心用例清单 ----------
rows = [
    ('1', 'DefaultScoreCalculator.score', '正常',
     '对子10+10无功能牌：基础分10、点数20、得分(10+20)×2=60', '否'),
    ('2', 'DefaultScoreCalculator.score', '正常',
     '奸诈小丑（对子+50积分）整手只触发一次：bonusChips=50，总分(10+20+50)×2=160', '否'),
    ('3', 'DefaultScoreCalculator.score', '边界',
     '高牌只计最大一张：A计11分，得分(5+11)×1=16；人头牌J/Q/K计10分', '否'),
    ('4', 'DefaultScoreCalculator.score', '正常',
     '成长类功能牌每打出一次永久+50积分：两次出牌后bonusChips=100，总分260', '否'),
    ('5', 'DefaultHandTypeEvaluator.evaluateDetail', '正常',
     '同花顺判定：5张同花连牌判为STRAIGHT_FLUSH且5张全部参与计分', '否'),
    ('6', 'DefaultHandTypeEvaluator.isPlayable', '异常',
     '选0张或6张不合法：isPlayable返回false，evaluateDetail返回empty', '否'),
    ('7', 'DefaultGameSession.play / outcome', '边界',
     '出牌次数耗尽且未达标判FAILED；达标判LEVEL_CLEARED；第8关达标判VICTORY', '否'),
]
t = d.tables[0]
for i, r in enumerate(rows, start=1):
    for j, v in enumerate(r):
        t.rows[i].cells[j].text = v

# ---------- 三、核心测试代码 ----------
def set_code_font(p):
    for run in p.runs:
        run.font.name = 'Times New Roman'
        run.font.size = Pt(9)  # 小五
        run._element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')

def insert_code_before(anchor_p, code):
    """在 anchor 段落前逐行插入代码，返回无"""
    lines = code.rstrip('\n').split('\n')
    for line in lines:
        np = anchor_p.insert_paragraph_before(line)
        set_code_font(np)

def find_para(text):
    for p in d.paragraphs:
        if p.text.strip().startswith(text):
            return p
    return None

def fill_class_name(prefix, name):
    p = find_para(prefix)
    p.text = p.text.replace('________', name)

# 读取测试类源码
with open('src/test/java/com/bilitro/model/hand/DefaultScoreCalculatorTest.java', encoding='utf-8') as f:
    code1 = f.read()
with open('src/test/java/com/bilitro/model/game/DefaultGameSessionTest.java', encoding='utf-8') as f:
    code2 = f.read()

fill_class_name('3.1 Model层核心测试类', 'DefaultScoreCalculatorTest')
fill_class_name('3.2 Model层核心测试类', 'DefaultGameSessionTest')

# 3.1：删掉占位说明，在 3.2 标题前插入代码
p_ph1 = find_para('（此处粘贴完整测试类代码')
p_32 = find_para('3.2 Model层核心测试类')
insert_code_before(p_32, code1)
p_ph1.text = ''

# 3.2：在 3.3 标题前插入代码
p_ph2 = find_para('（此处粘贴完整测试类代码')
p_33 = find_para('3.3 Controller层核心测试类')
insert_code_before(p_33, code2)
p_ph2.text = ''

# 3.3：本组无 Controller 层测试，写明原因
p_33.text = '3.3 Controller层核心测试类（类名：无）'
p_ph3 = find_para('（此处粘贴完整测试类代码')
p_ph3.text = ('本组未编写 Controller 层单元测试。Controller 依赖 JavaFX 界面交互（动画回调、弹窗、场景切换），'
              '单元测试需启动 UI 工具包，成本高收益低；本组测试策略为模型层追求高覆盖（JaCoCo 分支覆盖率 55%~88%），'
              'Controller/View 层通过人工试玩验收。')

d.save(PATH)
print('done')
