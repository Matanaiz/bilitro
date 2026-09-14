# -*- coding: utf-8 -*-
"""生成实训日志 docx（Day01-Day06）"""
from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

logs = [
    {
        "day": "Day01", "stage": "项目启动与工程环境",
        "work": [
            "确定选题：基于 JavaFX 的小丑牌（Balatro 风格）Roguelike 卡牌游戏，确认选题表与组内分工",
            "搭建 Maven 工程骨架，规划 model / view / controller 分层包结构，配置 JavaFX 17 与 JUnit 5 依赖",
            "创建 GitHub 远程仓库，完成首次代码推送，编写组员协作流程文档（克隆、分支、提交、拉取、推送）",
            "排查本机 Git 凭据与网络连接问题，确定 HTTPS + 凭据管理器的协作方案",
        ],
        "result": "选题确认表与分工计划（含 README）、Maven 工程骨架（可编译的空壳工程）",
        "problem": "命令行 Git 推送时反复要求身份验证且连接不稳定，改用凭据管理器并重试后推送成功",
        "next": "开展需求分析，梳理核心玩法与功能清单",
    },
    {
        "day": "Day02", "stage": "需求分析",
        "work": [
            "梳理核心玩法需求：出牌计分、牌型判定、功能牌（小丑牌）效果、商店购买与出售、关卡推进",
            "确定牌型优先级表及对应初始分数与倍数，确定功能牌触发规则（每张手牌触发一轮全部功能牌）",
            "确定关卡目标分数增长规则、代币与商店刷新费用规则、利息与过关奖励规则",
            "编写需求规格说明书并组内评审",
        ],
        "result": "需求规格说明书",
        "problem": "功能牌触发次序存在歧义（每张牌触发一轮还是整手触发一次），经讨论确定为每张手牌触发一轮，结算型功能牌（如小丑）在计分结束时只触发一次",
        "next": "进行概要设计，确定接口与包结构",
    },
    {
        "day": "Day03", "stage": "概要设计",
        "work": [
            "设计整体包结构与分层：model（card / hand / game / shop / player）、view、controller",
            "定义核心接口：ScoreCalculator、HandTypeEvaluator、GameSession、Shop 等，明确各接口职责",
            "绘制玩家回合内游玩时序图，补充游戏状态机（含暂停、商店、结算等状态）",
            "编写概要设计说明书并组内评审",
        ],
        "result": "概要设计说明书（含包结构图、接口定义、时序图、状态机图）",
        "problem": "结束判定与计分流程对接口有新增要求，根据评审讨论补充了回合结果回调与状态补充",
        "next": "完成详细设计与编码规范自查",
    },
    {
        "day": "Day04", "stage": "详细设计与编码规范",
        "work": [
            "编写详细设计说明书：类的字段、方法签名、计分流程伪代码、功能牌配置表设计",
            "确定功能牌以配置表方式实现，功能牌池清单预留占位，后续补充牌面图片与效果文本",
            "确定界面交互方案：手牌与功能牌采用按钮控件，手牌选中上移并高亮边框，功能牌点击缩放并弹出说明",
            "完成编码规范自查表（命名、注释、方法长度拆分等）",
        ],
        "result": "详细设计说明书、编码规范自查表",
        "problem": "功能牌上限与槽位展示方式经过讨论确定为六个固定卡槽居中排列，商店与主界面样式保持一致",
        "next": "进入核心编码阶段，实现主循环与核心玩法",
    },
    {
        "day": "Day05", "stage": "核心编码（主循环）",
        "work": [
            "实现核心玩法主循环：发牌、选牌（最多 5 张）、出牌计分、弃牌、关卡推进与结束判定",
            "实现计分过程动画：手牌从左到右逐张触发，当前分数与倍数随计算逐步增长并带缩放动效",
            "实现商店模块：购买、出售（半价）、刷新（费用递增、每关重置）、功能牌不重复出现",
            "实现功能牌二十余张，含重新触发型与成长型两类可复用接口，选中手牌时区分计分牌与不计分牌",
            "界面采用 BorderPane 布局，左侧显示当前分数与倍数，底部手牌按点数从大到小排列",
        ],
        "result": "游戏源码（核心玩法可运行，含商店与功能牌系统）",
        "problem": "通关弹窗在动画回调中直接弹出导致 showAndWait 异常，改为 Platform.runLater 延迟弹出后解决",
        "next": "编写单元测试，准备中期检查",
    },
    {
        "day": "Day06", "stage": "单元测试与中期检查",
        "work": [
            "为 model 层编写单元测试：牌型判定、计分器、游戏会话、商店随机逻辑，共 67 个用例全部通过",
            "配置 JaCoCo 覆盖率插件，解决中文路径导致执行数据文件缺失的问题，生成 HTML 覆盖率报告",
            "model 层行覆盖率约 77% 至 96%，view 与 controller 为界面层代码，按计划不纳入单元测试",
            "编写单元测试报告，整理项目文档与目录结构，更新风险清单",
        ],
        "result": "单元测试代码与运行输出、源码定稿、单元测试报告",
        "problem": "IDEA 未识别 JUnit 依赖导致编译报错，通过 Maven 刷新与清除缓存重启后解决；JaCoCo 报告路径含中文导致数据文件写入失败，改到用户目录后正常",
        "next": "根据中期检查意见完善功能，补充界面与文档",
    },
]

doc = Document()
style = doc.styles["Normal"]
style.font.name = "宋体"
style.font.size = Pt(11)

title = doc.add_heading("实训日志", level=0)
title.alignment = WD_ALIGN_PARAGRAPH.CENTER

for log in logs:
    doc.add_heading(f"{log['day']}  {log['stage']}", level=1)

    doc.add_paragraph("当日工作内容：", style=None).runs[0].bold = True
    for item in log["work"]:
        doc.add_paragraph(item, style="List Number")

    p = doc.add_paragraph()
    p.add_run("提交成果：").bold = True
    p.add_run(log["result"])

    p = doc.add_paragraph()
    p.add_run("遇到的问题与解决：").bold = True
    p.add_run(log["problem"])

    p = doc.add_paragraph()
    p.add_run("评审意见：").bold = True
    p.add_run("（待评审后填写）")

    p = doc.add_paragraph()
    p.add_run("次日计划：").bold = True
    p.add_run(log["next"])

out = r"D:\桌面\code\javafx\bilitro\docs\实训日志_Day01-Day06.docx"
doc.save(out)
print("saved:", out)
