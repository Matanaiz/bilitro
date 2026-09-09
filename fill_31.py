# -*- coding: utf-8 -*-
"""填写概要设计说明书 3.1 玩家回合内游玩流程，并插入时序图。"""
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

PATH = r"D:\桌面\code\javafx\bilitro\概要设计说明书-模板1.1.docx"
IMG = r"D:\桌面\code\javafx\bilitro\玩家回合内游玩时序图.drawio.png"
d = Document(PATH)

def find_para(text_start):
    for p in d.paragraphs:
        if p.text.strip().startswith(text_start):
            return p
    raise RuntimeError("not found: " + text_start)

def insert_before(anchor, text, bold=False):
    p = anchor.insert_paragraph_before("")
    run = p.add_run(text)
    run.font.size = Pt(10.5)
    run.font.bold = bold
    return p

# 标题改名
h31 = find_para("3.1 XXX业务流程")
h31.runs[0].text = "3.1 玩家回合内游玩流程"

# 3.1.1 场景说明
p311 = find_para("简单描述该时序图对应的业务场景")
p311.text = ""
h312 = find_para("3.1.2 交互流程")
scene = ("本时序图对应玩家进入对局后进行一个回合内操作的场景。参与对象包括玩家、"
         "游玩界面系统（视图层与控制层，负责界面展示、接收点击、转发用户意图）和"
         "游戏处理系统（模型层，负责数据初始化、弃牌条件检查、牌型判定与计分）。"
         "场景从玩家点击进入对局开始，覆盖选牌、弃牌、出牌计分三个核心操作，"
         "对应的时序图如下：")
p311.add_run(scene).font.size = Pt(10.5)
# 插图（放在 3.1.2 标题之前）
pic_para = h312.insert_paragraph_before("")
pic_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
pic_para.add_run().add_picture(IMG, width=Inches(4.5))
cap = h312.insert_paragraph_before("")
cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
cr = cap.add_run("图 3-1 玩家回合内游玩时序图")
cr.font.size = Pt(9)

# 3.1.2 交互流程
anchor312 = find_para("按顺序梳理完整调用、交互步骤、触发条件、返回结果。")
anchor312.text = ""
anchor32 = find_para("3.2 XXX业务流程")
steps = [
    "1. 玩家点击进入对局，游玩界面系统向游戏处理系统请求初始化数据。",
    "2. 游戏处理系统初始化本局数据（关卡目标分、出牌与弃牌次数、洗牌并发手牌），将初始数据返回给游玩界面系统。",
    "3. 游玩界面系统初始化界面，向玩家展示初始数据并显示手牌。",
    "4. 玩家点击选择手牌，游玩界面系统对选中的手牌给出高亮反馈（上移并显示选中边框）。",
    "5. 可选操作：玩家点击弃牌，游玩界面系统向游戏处理系统申请弃牌。",
    "6. 游戏处理系统检查弃牌条件（剩余弃牌次数是否大于零）：条件成立则进行弃牌处理，移除选中手牌并补发等量新牌；条件不成立则不进行弃牌。",
    "7. 游玩界面系统向玩家展示弃牌结果，刷新手牌显示。",
    "8. 玩家点击出牌，游玩界面系统将选中的手牌提交给游戏处理系统。",
    "9. 游戏处理系统判定牌型并计算功能牌效果：找出优先级最高的牌型，以其基础分和基础倍数初始化，在牌型对应的手牌中从左往右逐张累加点数，每处理一张手牌触发一轮功能牌，最终得分等于最终分数乘最终倍数。",
    "10. 游戏处理系统将本次得分返回给游玩界面系统。",
    "11. 游玩界面系统更新对局总分与剩余目标分，刷新界面显示，等待玩家下一步操作。",
]
for s in steps:
    insert_before(anchor32, s)

d.save(PATH)
print("saved")
